# Copyright (c) 2026, sammish and contributors
# For license information, please see license.txt

"""Apply HR's employee master workbook: name, title, department/section, shift and ID card.

Rows are the "All Employees" sheet: E No., Employee Name, Title, Department, Section, Shift,
ID Card No. Employees are matched on E No. (the TAMS code, kept as employee_number); the
workbook is the current master, so its names and titles replace the TAMS ones. Safe to re-run.
"""

import frappe
from frappe.utils import cstr, nowdate

# workbook shift -> Shift Type created for it
SHIFTS = {
	"Flexible (8 hrs/day)": {
		"name": "Flexible 8 Hours",
		"start_time": "05:00:00",
		"end_time": "04:59:00",
		"custom_is_flexible_shift": 1,
		"custom_day_start_time": "05:00:00",
		"custom_required_working_hours": 8,
	},
	"07:30 - 15:00 (Fixed)": {
		"name": "Fixed 07:30-15:00",
		"start_time": "07:30:00",
		"end_time": "15:00:00",
		"begin_check_in_before_shift_start_time": 60,
		# punches up to 21:00 still belong to the shift, so late work counts as overtime
		"allow_check_out_after_shift_end_time": 360,
	},
}
CONTRACTOR = "Contractor"


@frappe.whitelist(methods=["POST"])
def import_employee_master(rows, company=None):
	"""rows: [e_no, name, title, department, section, shift, id_card_no]"""
	frappe.only_for("System Manager")
	from raysut_cement.raysut_cement.tams_import import default_company

	company = company or default_company()
	rows = frappe.parse_json(rows)
	shifts = {label: ensure_shift_type(spec) for label, spec in SHIFTS.items()}

	by_number = {
		e.employee_number: e.name
		for e in frappe.get_all("Employee", {"employee_number": ("is", "set")}, ["name", "employee_number"])
	}
	unnumbered = {
		e.employee_name: e.name
		for e in frappe.get_all("Employee", {"employee_number": ("is", "not set"), "company": company}, ["name", "employee_name"])
	}
	result = {"updated": 0, "created": 0, "renamed": [], "errors": []}
	seen = set()

	for e_no, name, title, department, section, shift, id_card in rows:
		e_no, name = cstr(e_no).strip(), " ".join(cstr(name).split())
		if not name:
			continue
		existing = by_number.get(e_no) if e_no else unnumbered.get(name)
		employee = frappe.get_doc("Employee", existing) if existing else frappe.new_doc("Employee")
		if existing and employee.employee_name != name:
			result["renamed"].append([e_no, employee.employee_name, name])

		first, *middle, last = name.split() if len(name.split()) > 1 else (name, "")
		employee.update(
			{
				"first_name": first,
				"middle_name": " ".join(middle),
				"last_name": last,
				"company": company,
				"status": "Active",
				"designation": ensure_designation(title),
				"department": ensure_department(department, section, company),
				"default_shift": shifts.get(cstr(shift).strip()),
				"custom_id_card_no": cstr(id_card).strip() or None,
			}
		)
		if cstr(department).strip() == CONTRACTOR and frappe.db.exists("Employment Type", "Contract"):
			employee.employment_type = "Contract"
		if not existing and e_no:
			employee.employee_number = e_no
			if not frappe.db.exists("Employee", {"attendance_device_id": e_no}):
				employee.attendance_device_id = e_no  # TAMS convention: the clock card is the E No.

		# the workbook has no date of birth, gender or (for new people) joining date
		employee.flags.ignore_mandatory = True
		frappe.db.savepoint("employee_master")
		try:
			employee.save()
		except frappe.ValidationError as e:
			frappe.db.rollback(save_point="employee_master")
			result["errors"].append(f"{e_no or name}: {e}")
			continue
		seen.add(employee.name)
		result["updated" if existing else "created"] += 1

	result["active_not_in_master"] = frappe.get_all(
		"Employee", {"company": company, "status": "Active", "name": ("not in", list(seen) or [""])}, pluck="employee_number"
	)
	return result


def import_employee_master_file(path, company=None, sheet="All Employees"):
	"""bench execute entry point: read the workbook's sheet and apply it."""
	from openpyxl import load_workbook

	workbook = load_workbook(path, read_only=True, data_only=True)
	rows = [list(r[:7]) for r in list(workbook[sheet].iter_rows(values_only=True))[1:] if any(r)]
	result = import_employee_master(rows, company)
	frappe.db.commit()
	return result


def ensure_shift_type(spec):
	if not frappe.db.exists("Shift Type", spec["name"]):
		frappe.get_doc(
			{
				"doctype": "Shift Type",
				"__newname": spec["name"],
				"enable_auto_attendance": 1,
				"process_attendance_after": nowdate(),  # imported TAMS history is never re-marked
				"determine_check_in_and_check_out": "Alternating entries as IN and OUT during the same shift",
				"working_hours_calculation_based_on": "First Check-in and Last Check-out",
				**{k: v for k, v in spec.items() if k != "name"},
			}
		).insert()
	return spec["name"]


def ensure_designation(title):
	title = " ".join(cstr(title).split())
	if not title:
		return None
	# names match case-insensitively: reuse "FITTER" for "Fitter" under its own spelling
	return frappe.db.exists("Designation", title) or frappe.get_doc({"doctype": "Designation", "designation_name": title}).insert().name


def ensure_department(group, section, company):
	"""Section under its department group (Admin / Tech / Contractor); an existing one is reused as is."""
	group, section = " ".join(cstr(group).split()), " ".join(cstr(section).split())
	parent = find_department(group, company) if group else None
	if group and not parent:
		root = frappe.db.get_value("Department", {"parent_department": ("is", "not set"), "is_group": 1}) or None
		parent = (
			frappe.get_doc(
				{"doctype": "Department", "department_name": group, "company": company, "is_group": 1, "parent_department": root}
			)
			.insert()
			.name
		)
	if not section:
		return parent
	existing = find_department(section, company)
	if existing:
		return existing
	if parent and not frappe.db.get_value("Department", parent, "is_group"):
		parent = None  # an existing plain department named like the group: the section goes at the top
	return (
		frappe.get_doc({"doctype": "Department", "department_name": section, "company": company, "parent_department": parent})
		.insert()
		.name
	)


def find_department(department_name, company):
	return frappe.db.get_value("Department", {"department_name": department_name, "company": company}) or frappe.db.get_value(
		"Department", {"department_name": department_name, "company": ("is", "not set")}
	)
