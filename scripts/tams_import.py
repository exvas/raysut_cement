"""Send a TAMS export (scripts/tams_export.sh) to a site over the REST API, batch by batch.

	RC_TOKEN=api_key:api_secret python3 scripts/tams_import.py https://site OUT_DIR [--company NAME]
	RC_TOKEN=... python3 scripts/tams_import.py https://site --employee-master master.xlsx [--company NAME]

--employee-master applies HR's employee master workbook ("All Employees" sheet) after the TAMS
import: names, titles, departments/sections, shifts and ID card numbers.

The token's user needs System Manager. Every batch skips what the site already has, so an
interrupted run can simply be started again.
"""

import argparse
import csv
import json
import os
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

METHOD = "/api/method/raysut_cement.raysut_cement."


# the same three helpers as raysut_cement/raysut_cement/tams_import.py, which needs frappe to import
def read_tsv(path, filename):
	with open(os.path.join(path, filename), newline="", encoding="utf-8", errors="replace") as f:
		return [[value.strip() for value in row] for row in csv.reader(f, delimiter="\t", quoting=csv.QUOTE_NONE)]


def batches(rows, size):
	for start in range(0, len(rows), size):
		yield rows[start : start + size]


def employee_batches(rows, size):
	batch = []
	for row in rows:
		if len(batch) >= size and row[0] != batch[-1][0]:
			yield batch
			batch = []
		batch.append(row)
	if batch:
		yield batch


def call(site, method, **kwargs):
	"""method: a tams_import function, or module.function under raysut_cement.raysut_cement"""
	method = method if "." in method else f"tams_import.{method}"
	request = urllib.request.Request(
		site.rstrip("/") + METHOD + method,
		data=json.dumps(kwargs).encode(),
		headers={"Authorization": f"token {os.environ['RC_TOKEN']}", "Content-Type": "application/json"},
	)
	for attempt in range(5):
		try:
			with urllib.request.urlopen(request, timeout=600) as response:
				return json.load(response)["message"]
		except urllib.error.HTTPError as e:
			if e.code < 500 or attempt == 4:
				sys.exit(f"{method}: HTTP {e.code} {e.read().decode()[:2000]}")
		except (urllib.error.URLError, TimeoutError) as e:
			if attempt == 4:
				raise
			print(f"  retrying after {e}")
		time.sleep(5 * (attempt + 1))


def active_employee_numbers(site, company=None):
	filters = [["status", "=", "Active"]] + ([["company", "=", company]] if company else [])
	query = urllib.parse.urlencode(
		{"filters": json.dumps(filters), "fields": json.dumps(["employee_number"]), "limit_page_length": 0}
	)
	request = urllib.request.Request(
		f"{site.rstrip('/')}/api/resource/Employee?{query}", headers={"Authorization": f"token {os.environ['RC_TOKEN']}"}
	)
	with urllib.request.urlopen(request, timeout=120) as response:
		return [row["employee_number"] for row in json.load(response)["data"]]


def main():
	parser = argparse.ArgumentParser()
	parser.add_argument("site")
	parser.add_argument("path", nargs="?")
	parser.add_argument("--company")
	parser.add_argument("--only", choices=["employees", "checkins", "timesheet"])
	parser.add_argument("--employee-master", metavar="XLSX")
	args = parser.parse_args()

	if args.employee_master:
		from openpyxl import load_workbook

		sheet = load_workbook(args.employee_master, read_only=True, data_only=True)["All Employees"]
		rows = [list(r[:7]) for r in list(sheet.iter_rows(values_only=True))[1:] if any(r)]
		totals = {"updated": 0, "created": 0, "renamed": 0}
		for i, batch in enumerate(batches(rows, 25)):
			result = call(args.site, "employee_master.import_employee_master", rows=batch, company=args.company)
			for key in totals:
				totals[key] += len(result[key]) if key == "renamed" else result[key]
			for error in result["errors"]:
				print("  error:", error)
			print(f"employee master {min((i + 1) * 25, len(rows))}/{len(rows)}: {totals}", flush=True)
		in_master = {str(r[0]) for r in rows if r[0] is not None}
		active = active_employee_numbers(args.site, args.company)
		print("active employees not in the workbook:", len([n for n in active if n not in in_master]))
		return

	if args.only in (None, "employees"):
		# about a second per employee on the server: small batches stay inside the gateway timeout
		rows, created = read_tsv(args.path, "employees.tsv"), 0
		for i, batch in enumerate(batches(rows, 25)):
			result = call(args.site, "import_employees", rows=batch, company=args.company)
			created += result["created"]
			for error in result["errors"]:
				print("  error:", error)
			print(f"employees {min((i + 1) * 25, len(rows))}/{len(rows)} read, {created} created", flush=True)

	if args.only in (None, "checkins"):
		rows = sorted(read_tsv(args.path, "checkins.tsv"), key=lambda row: row[1])
		total = 0
		for i, batch in enumerate(batches(rows, 2000)):
			total += call(args.site, "import_checkins", rows=batch)["created"]
			print(f"checkins {min((i + 1) * 2000, len(rows))}/{len(rows)} read, {total} created", flush=True)

	if args.only in (None, "timesheet"):
		rows = read_tsv(args.path, "timesheet.tsv")
		total = done = 0
		for batch in employee_batches(rows, 3000):
			total += call(args.site, "import_timesheet", rows=batch, company=args.company)["created"]
			done += len(batch)
			print(f"timesheet {done}/{len(rows)} read, {total} attendance created", flush=True)


if __name__ == "__main__":
	main()
